#!/usr/bin/env python3
"""Turn an existing PowerPoint template into an mdslide master.

Usage:
    python tools/convert_master.py template.pptx -o master.pptx [--profile report|presentation]

The template itself is the base and nothing is copied between packages, so its theme, fonts, colours, logos, header
and footer stay exactly where they are. Two layers are separated on purpose (ADR-0030, ADR-0033):

  意匠   theme, fonts, colours, background, decorations, header, footer, Cover and Section — left untouched.
  版面   the body pages (Body-Text / Agenda / Body-2col): their margins, title band, body box, gutter, type sizes
         and bullet indent come from the bundled sample masters (tools/master_profiles.py).

The second layer is what makes "one page holds N lines" a property of mdslide instead of whichever template someone
brought: mdslide measures the body from the master (src/model/boxes.ts), so a template whose body style is 32pt
would otherwise fit seven lines to a page and split everything else onto continuations.
"""
from __future__ import annotations

import argparse
import copy
import re
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn

from master_profiles import (DEFAULT_PROFILE, EMU_PER_PT, HANG_EM, INSET_LR, INSET_TB, PROFILES, Profile)

SAFE_GAP = 0.01    # mdslide's own clearance from the header and the footer (src/model/boxes.ts)
EDGE_ZONE = 1 / 3  # a shape ending in the top third is a header; one starting in the bottom third is a footer
EMU_PER_CM = 360000
TITLES = ("title", "ctrTitle")
BODIES = ("body", "obj")
FURNITURE = ("dt", "ftr", "sldNum")  # date / footer / slide number: kept, the exporter copies them onto slides
ROLES = ["Cover", "Body-2col", "Section", "Body-Text", "Agenda"]  # assignment order: the telling ones first
BODY_PAGES = ("Body-Text", "Agenda", "Body-2col")  # the only roles whose measure and type are replaced
PH_PATH = f'{qn("p:nvSpPr")}/{qn("p:nvPr")}/{qn("p:ph")}'
# Source names that usually hold a role, matched on the name lowercased with spaces removed.
HINTS = {
    "Cover": ("cover", "titleslide", "表紙", "とびら"),
    "Agenda": ("agenda", "toc", "目次", "contents", "titleonly", "もくじ"),
    "Section": ("section", "chapter", "divider", "中表紙", "章", "扉"),
    "Body-Text": ("bodytext", "titleandcontent", "content", "本文"),
    "Body-2col": ("body2col", "twocontent", "comparison", "2col", "twocol", "二段", "2段", "比較"),
}
PROFILE: Profile = PROFILES[DEFAULT_PROFILE]  # main() replaces it; helpers called on their own still have a measure


def reads_as_role(name: str) -> str | None:
    """The role mdslide reads from a layout name. Mirror of roleFromLayoutName (src/master/importMaster.ts)."""
    n = "".join(name.split()).lower()
    if n.startswith(("cover", "titleslide", "title-slide", "表紙")):
        return "Cover"
    if n.startswith(("agenda", "toc", "目次")):
        return "Agenda"
    if n.startswith(("section", "chapter", "中表紙", "章")):
        return "Section"
    if n.startswith("body"):
        rest = n[4:]
        return {"text": "Body-Text", "2col": "Body-2col", "twocol": "Body-2col"}.get(rest[1:] if rest.startswith("-") else rest)
    return None


def kind(shape) -> str | None:
    """The placeholder type as the OOXML spells it ("title", "body", "sldNum", ...), or None for a plain shape."""
    el = shape._element.find(PH_PATH)
    return None if el is None else (el.get("type") or "obj")


def rect(shape):
    """Effective position in EMU (python-pptx resolves what the layout inherits), or None when it has none."""
    try:
        r = (shape.left, shape.top, shape.width, shape.height)
    except (AttributeError, TypeError, ValueError):
        return None
    return None if any(v is None for v in r) else r


def by_area(shapes):
    """Largest first, document order breaking ties: the rule bodyPlaceholders and body_placeholders already use."""
    sized = []
    for i, shape in enumerate(shapes):
        r = rect(shape)
        if r:
            sized.append((-(r[2] * r[3]), i, shape))
    return [shape for _, _, shape in sorted(sized, key=lambda t: t[:2])]


def placeholders(layout, kinds):
    return [s for s in layout.placeholders if kind(s) in kinds]


def band(layout, slide_w, slide_h):
    """Bottom of the header and top of the footer in EMU, from the master's and the layout's decorations and the
    date / footer / slide-number placeholders. Mirror of contentBand (src/model/boxes.ts)."""
    shapes = [] if layout.element.get("showMasterSp") == "0" else [s for s in layout.slide_master.shapes if not s.is_placeholder]
    shapes += [s for s in layout.shapes if not s.is_placeholder]
    shapes += placeholders(layout, FURNITURE)
    left, right = PROFILE.margin * slide_w, (1 - PROFILE.margin) * slide_w
    header = footer = None
    for s in shapes:
        r = rect(s)
        if not r:
            continue
        x, y, w, h = r
        if w <= 0 or h <= 0 or x + w <= left or x >= right:  # a side stripe never blocks the content columns
            continue
        if y + h <= slide_h * EDGE_ZONE:
            header = max(header or 0, y + h)
        elif y >= slide_h * (1 - EDGE_ZONE):
            footer = min(slide_h if footer is None else footer, y)
    return header, footer


def title_box(slide_w, slide_h):
    """The profile's title band: where a body page's title goes."""
    margin = PROFILE.margin * slide_w
    return (int(margin), int(margin), int(slide_w - 2 * margin), int(PROFILE.title_band * slide_h))


def content_box(layout, slide_w, slide_h):
    """Where the body goes: full width inside the margins, under the title, clear of the header and the footer.

    Only a title in the top third counts. Office's caption layouts put theirs at 70% of the height, and starting the
    body under such a title leaves a three-line strip above the footer; normalize moves those titles up instead."""
    titles = [r for r in (rect(s) for s in placeholders(layout, TITLES)) if r and r[1] < slide_h * EDGE_ZONE]
    default_bottom = PROFILE.margin * slide_w + PROFILE.title_band * slide_h
    title_bottom = max((r[1] + r[3] for r in titles), default=int(default_bottom))
    header, footer = band(layout, slide_w, slide_h)
    top = max(title_bottom + PROFILE.title_gap * slide_w, 0 if header is None else header + SAFE_GAP * slide_w)
    bottom = min(slide_h - PROFILE.margin * slide_w, slide_h if footer is None else footer - SAFE_GAP * slide_w)
    return (int(PROFILE.margin * slide_w), int(top), int(slide_w - 2 * PROFILE.margin * slide_w), int(max(0.01 * slide_w, bottom - top)))


def columns(box, slide_w):
    """The two column boxes of a 2col layout, left to right, the profile's gutter apart."""
    x, y, w, h = box
    half = (w - int(PROFILE.gutter * slide_w)) // 2
    return (x, y, half, h), (x + w - half, y, half, h)


def drop(shape):
    shape._element.getparent().remove(shape._element)


def place(shape, box):
    shape.left, shape.top, shape.width, shape.height = box


def list_style(shape):
    """The placeholder's own a:lstStyle, created right after a:bodyPr when it has none.

    The order inside a:txBody is bodyPr, lstStyle, then the paragraphs: appending would produce a file PowerPoint
    refuses to open."""
    body = shape._element.txBody
    lst = body.find(qn("a:lstStyle"))
    if lst is None:
        lst = body.makeelement(qn("a:lstStyle"), {})
        pr = body.find(qn("a:bodyPr"))
        body.insert(list(body).index(pr) + 1 if pr is not None else 0, lst)
    return lst


LEVEL_TAG = re.compile(r"\}lvl(\d)pPr$")


def level(lst, number):
    """One a:lvlNpPr of a list style, inserted in ascending order when it is missing.

    The levels must stay in order inside a:lstStyle, so a missing lvl1 goes before an existing lvl2."""
    tag = qn(f"a:lvl{number}pPr")
    found = lst.find(tag)
    if found is not None:
        return found
    lvl = lst.makeelement(tag, {})
    later = [e for e in lst for m in [LEVEL_TAG.search(str(e.tag))] if m and int(m.group(1)) > number]
    lst.insert(list(lst).index(later[0]) if later else len(lst), lvl)
    return lvl


def set_type(shape, sizes, hang=False):
    """Write the profile's type into one placeholder's own list style and reset its text insets.

    The layout's lstStyle wins over the master's txStyles, so this is both how the size reaches PowerPoint and
    mdslide (importMaster reads the layout first) and how a size the source template left behind is overridden."""
    lst = list_style(shape)
    for number, size in enumerate(sizes, start=1):
        lvl = level(lst, number)
        if hang:
            step = round(size * HANG_EM * EMU_PER_PT)
            lvl.set("marL", str(step * number))
            lvl.set("indent", str(-step))
        rpr = lvl.find(qn("a:defRPr"))
        if rpr is None:
            rpr = etree.SubElement(lvl, qn("a:defRPr"))
        rpr.set("sz", str(round(size * 100)))
    pr = shape._element.txBody.find(qn("a:bodyPr"))
    if pr is not None:
        for name, value in (("lIns", INSET_LR), ("rIns", INSET_LR), ("tIns", INSET_TB), ("bIns", INSET_TB)):
            pr.set(name, str(value))


def clone_layout(prs, base, name):
    """Copy a layout part inside the same package (rels included) and register it on its master."""
    pkg = prs.part.package
    part = type(base.part).load(pkg.next_partname("/ppt/slideLayouts/slideLayout%d.xml"), base.part.content_type, pkg, base.part.blob)
    for rel in base.part.rels.values():
        if rel.is_external:
            part.rels.get_or_add_ext_rel(rel.reltype, rel.target_ref)
        else:
            part.relate_to(rel.target_part, rel.reltype)
    master = base.slide_master
    rid = master.part.relate_to(part, RT.SLIDE_LAYOUT)
    lst = master.element.find(qn("p:sldLayoutIdLst"))
    ids = [int(e.get("id")) for e in lst.findall(qn("p:sldLayoutId"))] or [2147483648]
    lst.append(lst.makeelement(qn("p:sldLayoutId"), {"id": str(max(ids) + 1), qn("r:id"): rid}))
    layout = part.slide_layout
    layout.name = name
    return layout


def clone_placeholder(layout, source, ph_type, box):
    """Duplicate a placeholder inside a layout (same package, so no relationships to carry) as another type."""
    el = copy.deepcopy(source._element)
    tree = layout.shapes._spTree
    next_id = max(int(e.get("id")) for e in tree.iter(qn("p:cNvPr"))) + 1
    next_idx = max((int(s._element.find(PH_PATH).get("idx") or 0) for s in layout.placeholders), default=0) + 1
    el.find(f'{qn("p:nvSpPr")}/{qn("p:cNvPr")}').set("id", str(next_id))
    ph = el.find(PH_PATH)
    ph.set("type", ph_type)
    ph.set("idx", str(next_idx))
    tree.append(el)
    shape = next(s for s in layout.placeholders if s.shape_id == next_id)
    place(shape, box)
    return shape


def score(role, layout):
    """How well a source layout suits a role: its name first, then the placeholders it carries.

    A page built around a picture is never a text page, however few candidates are left: Office's "Picture with
    Caption" carries a title and one caption box, which used to be enough to win the agenda."""
    n = "".join(layout.name.split()).lower()
    bodies, titles = placeholders(layout, BODIES), placeholders(layout, TITLES)
    slide_h = layout.slide_master.part.package.presentation_part.presentation.slide_height
    pts = 0
    if reads_as_role(layout.name) == role:
        pts += 100
    if any(h in n for h in HINTS[role]):
        pts += 50
    if role == "Cover":
        pts += 20 if placeholders(layout, ("subTitle",)) or any(kind(s) == "ctrTitle" for s in titles) else 0
    else:
        if placeholders(layout, ("pic",)):
            pts -= 40                                    # a caption page, not a text page
        top = min((r[1] for r in (rect(s) for s in titles) if r), default=0)
        if top >= slide_h * EDGE_ZONE:
            pts -= 30                                    # the title sits where mdslide would never put one
        if role == "Body-2col":
            pts += 30 if len(bodies) >= 2 else -30
        else:                                            # a spare title-only layout is fine: normalize adds the box
            pts += 15 if len(bodies) <= 1 else 0
    return pts + (5 if titles else 0)


def pick_roles(layouts):
    """One source layout per role, best match first; a role with no convincing candidate is left out (it is cloned)."""
    chosen, taken = {}, set()
    for role in ROLES:
        ranked = sorted(((score(role, l), i, l) for i, l in enumerate(layouts) if id(l) not in taken), key=lambda t: (-t[0], t[1]))
        if ranked and ranked[0][0] >= 20:
            chosen[role] = ranked[0][2]
            taken.add(id(ranked[0][2]))
    return chosen


def normalize(layout, role, slide_w, slide_h):
    """Give a role layout exactly the boxes mdslide fills.

    Decorations, theme and furniture are left untouched everywhere. On the body pages the title band, the body box
    and the type come from the profile; the cover and the section keep the size and the position the designer chose."""
    titles = by_area(placeholders(layout, TITLES))
    if not titles and not placeholders(layout, BODIES):
        raise SystemExit(f"レイアウト '{layout.name}' に枠がありません。別のテンプレートを渡してください。")
    keep = [s.shape_id for s in titles[:1]]
    if role in BODY_PAGES and titles:
        place(titles[0], title_box(slide_w, slide_h))
        set_type(titles[0], (PROFILE.title_pt,))
    box = content_box(layout, slide_w, slide_h)
    if role == "Cover":  # the cover is the most bespoke page: leave its boxes and its type where the designer put them
        subs = by_area(placeholders(layout, ("subTitle",))) or by_area(placeholders(layout, BODIES))
        if not subs:
            subs = [clone_placeholder(layout, titles[0], "subTitle", (box[0], box[1], box[2], box[3] // 3))]
        keep += [s.shape_id for s in subs[:1]]
    elif role != "Section":  # a section page carries a title and nothing mdslide would fill
        wanted = 2 if role == "Body-2col" else 1
        found = by_area(placeholders(layout, BODIES))
        boxes = columns(box, slide_w) if wanted == 2 else (box,)
        while len(found) < wanted:
            found.append(clone_placeholder(layout, found[0] if found else titles[0], "body", boxes[len(found)]))
        found = sorted(found[:wanted], key=lambda s: rect(s)[0]) if wanted == 2 else found[:1]
        for shape, target in zip(found, boxes):
            place(shape, target)
            set_type(shape, PROFILE.body_pts, hang=True)
        keep += [s.shape_id for s in found]
    for shape in list(layout.placeholders):
        if shape.shape_id not in keep and kind(shape) not in FURNITURE:
            drop(shape)


def describe(layout):
    """One report line per layout: what mdslide will find in it."""
    parts = []
    for shape in layout.placeholders:
        r = rect(shape)
        size = f"{r[2] / EMU_PER_CM:.1f}x{r[3] / EMU_PER_CM:.1f}cm" if r else "継承"
        parts.append(f"{kind(shape)}({shape.placeholder_format.idx}) {size}")
    body = by_area(placeholders(layout, BODIES))
    lines = ""
    if body:
        height_pt = rect(body[0])[3] / EMU_PER_PT
        lines = f"  本文およそ {max(3, int(height_pt / (PROFILE.body_pts[0] * 1.3)))} 行"
    return f"  {layout.name:<10} {', '.join(parts)}{lines}"


def main():
    global PROFILE
    ap = argparse.ArgumentParser(description="手持ちの PowerPoint テンプレートを mdslide のマスターに変換します。")
    ap.add_argument("source", help="元のテンプレート (.pptx / .potx)")
    ap.add_argument("-o", "--out", required=True, help="書き出すマスター (master.pptx)")
    ap.add_argument("--profile", choices=sorted(PROFILES), default=DEFAULT_PROFILE,
                    help=f"本文ページの版面と文字の基準にする見本（既定: {DEFAULT_PROFILE}）")
    args = ap.parse_args()
    PROFILE = PROFILES[args.profile]
    src, out = Path(args.source), Path(args.out)
    if not src.exists():
        raise SystemExit(f"元のファイルがありません: {src}")

    prs = Presentation(str(src))
    slide_w, slide_h = prs.slide_width, prs.slide_height

    sld_ids = prs.slides._sldIdLst
    dropped = len(sld_ids)
    for sld_id in list(sld_ids):
        prs.part.drop_rel(sld_id.rId)
        sld_ids.remove(sld_id)

    layouts = list(prs.slide_layouts)
    if not layouts:
        raise SystemExit("スライドマスターにレイアウトがありません。PowerPoint で開けるテンプレートを渡してください。")
    chosen = pick_roles(layouts)
    base = chosen.get("Body-Text") or chosen.get("Agenda") or chosen.get("Section") or layouts[0]
    origin = {role: layout.name for role, layout in chosen.items()}
    for role in ROLES:
        if role not in chosen:
            chosen[role] = clone_layout(prs, base, role)
            origin[role] = f"{base.name} の複製"

    # Layouts mdslide would also read as a role would shadow the chosen ones (the first match wins), so move them aside.
    picked = {layout.part.partname for layout in chosen.values()}
    for master in prs.slide_masters:
        for layout in master.slide_layouts:
            if layout.part.partname not in picked and reads_as_role(layout.name):
                layout.name = f"Unused-{layout.name}"
    for role, layout in chosen.items():
        layout.name = role
        normalize(layout, role, slide_w, slide_h)

    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))

    check = Presentation(str(out))
    found = {reads_as_role(l.name) for l in check.slide_layouts}
    missing = [r for r in ("Cover", "Agenda", "Section", "Body-Text") if r not in found]
    if missing:
        raise SystemExit(f"変換に失敗しました。役割が足りません: {', '.join(missing)}")
    print(f"{out} を書き出しました（{slide_w / EMU_PER_CM:.1f}x{slide_h / EMU_PER_CM:.1f}cm、"
          f"見本「{PROFILE.name}」基準・本文 {PROFILE.body_pts[0]}pt、元のスライド {dropped} 枚を削除）")
    for role in ("Cover", "Agenda", "Section", "Body-Text", "Body-2col"):
        layout = check.slide_layouts.get_by_name(role)
        print(f"{role:<10} ← {origin.get(role, '')}")
        if layout:
            print(describe(layout))
    print("フォント・配色・ロゴ・ヘッダ・フッタ・表紙・中表紙は元のテンプレートのままです。"
          "本文ページの余白と文字だけ見本に合わせました。設定（⌘,）のマスターで取り込んでください。")


if __name__ == "__main__":
    main()
