"""Build examples/decorated-master.pptx: the sample master dressed up the way a corporate one is, for testing what the
preview and the export carry over from a master.

    python3 scripts/make_decorated_master.py          # needs python-pptx and Pillow (requirements.txt)

What it adds to examples/sample-master.pptx (layout names stay Cover / Agenda / Section / Body-Text / Body-2col):
  - theme: fonts Avenir Next / Hiragino Sans, a navy / teal / orange colour scheme
  - master: a header band with a dotted pattern and the logo, a footer band with an icon, fixed footer text, the date
    and the slide number; the title and body placeholders are moved below the band
  - Cover: a full-bleed generated hero image as the background, three picture tiles, a large logo, white title and
    subtitle at the bottom left; the master's bands are hidden (showMasterSp="0")
  - Section: a picture band down the left, the title moved to the right
All images are drawn with Pillow (gradients, shapes, silhouettes): nothing is downloaded, nothing is licensed."""
import copy
import math
import random
import tempfile
from pathlib import Path

from lxml import etree
from PIL import Image, ImageDraw, ImageFilter
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.oxml.shapes.autoshape import CT_Shape
from pptx.oxml.shapes.picture import CT_Picture
from pptx.shapes.autoshape import Shape
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "examples" / "sample-master.pptx"
OUT = ROOT / "examples" / "decorated-master.pptx"

NAVY, TEAL, ORANGE, INK, PAPER = (11, 61, 145), (0, 166, 166), (255, 122, 26), (27, 42, 65), (238, 242, 247)
FONT_LATIN, FONT_JA = "Avenir Next", "Hiragino Sans"

# ---- images (Pillow) -----------------------------------------------------------------------------------------------

def gradient(w, h, c1, c2):
    im = Image.new("RGB", (w, h))
    px = im.load()
    for y in range(h):
        for x in range(w):
            t = (x / w * 0.6 + y / h * 0.4)
            px[x, y] = tuple(round(a + (b - a) * t) for a, b in zip(c1, c2))
    return im


def hero(path):
    """Night skyline over a navy-to-teal sky: the cover's full-bleed background."""
    w, h = 1920, 1080
    im = gradient(w, h, NAVY, TEAL)
    glow = Image.new("RGBA", (w, h), (0, 0, 0, 0))
    g = ImageDraw.Draw(glow)
    rnd = random.Random(7)
    for _ in range(9):
        r = rnd.randint(120, 420); x, y = rnd.randint(0, w), rnd.randint(0, h // 2)
        g.ellipse((x - r, y - r, x + r, y + r), fill=(255, 255, 255, 22))
    glow = glow.filter(ImageFilter.GaussianBlur(40))
    im = Image.alpha_composite(im.convert("RGBA"), glow)
    d = ImageDraw.Draw(im)
    x = 0
    while x < w:  # skyline
        bw, bh = rnd.randint(40, 140), rnd.randint(120, 460)
        d.rectangle((x, h - bh, x + bw, h), fill=(12, 24, 48, 255))
        for wy in range(h - bh + 18, h - 12, 26):
            for wx in range(x + 8, x + bw - 10, 22):
                if rnd.random() < 0.35:
                    d.rectangle((wx, wy, wx + 8, wy + 12), fill=(255, 214, 120, 200))
        x += bw + rnd.randint(6, 30)
    streaks = Image.new("RGBA", (w, h), (0, 0, 0, 0))  # faint diagonal light streaks, composited so the alpha counts
    s = ImageDraw.Draw(streaks)
    for i in range(0, w, 160):
        s.line((i, 0, i + 600, h), fill=(255, 255, 255, 14), width=40)
    im = Image.alpha_composite(im, streaks)
    im.convert("RGB").save(path, optimize=True)


def logo(path, size=600):
    """Hexagon mark with an orange dot: the company logo."""
    im = Image.new("RGBA", (size, size), (0, 0, 0, 0))
    d = ImageDraw.Draw(im)
    cx = cy = size / 2; r = size * 0.46
    pts = [(cx + r * math.cos(math.radians(60 * i - 30)), cy + r * math.sin(math.radians(60 * i - 30))) for i in range(6)]
    d.polygon(pts, fill=NAVY + (255,))
    d.ellipse((cx - r * 0.55, cy - r * 0.55, cx + r * 0.55, cy + r * 0.55), fill=(255, 255, 255, 255))
    d.ellipse((cx - r * 0.2, cy - r * 0.2, cx + r * 0.2, cy + r * 0.2), fill=ORANGE + (255,))
    im.save(path, optimize=True)


def tile(path, kind):
    """Photo-like abstract tiles for the cover: mountains, waves, grid."""
    w, h = 800, 600
    if kind == "mountains":
        im = gradient(w, h, (255, 170, 90), (120, 60, 120))
        d = ImageDraw.Draw(im)
        rnd = random.Random(3)
        for layer, col in enumerate([(90, 70, 110), (60, 50, 90), (35, 30, 60)]):
            pts = [(0, h)]
            for x in range(0, w + 1, 80):
                pts.append((x, h * (0.45 + 0.12 * layer) + rnd.randint(-70, 70)))
            pts.append((w, h))
            d.polygon(pts, fill=col)
    elif kind == "waves":
        im = gradient(w, h, TEAL, NAVY)
        d = ImageDraw.Draw(im)
        for k in range(12):
            pts = [(x, h * 0.35 + k * 28 + 22 * math.sin(x / 60 + k)) for x in range(0, w + 1, 8)]
            d.line(pts, fill=(255, 255, 255, 255) if k % 3 == 0 else (200, 235, 240), width=3)
    else:
        im = gradient(w, h, PAPER, (200, 210, 225))
        d = ImageDraw.Draw(im)
        for x in range(0, w, 50):
            d.line((x, 0, x, h), fill=(170, 185, 205), width=1)
        for y in range(0, h, 50):
            d.line((0, y, w, y), fill=(170, 185, 205), width=1)
        rnd = random.Random(11)
        for _ in range(14):
            x, y = rnd.randrange(0, w, 50), rnd.randrange(0, h, 50)
            d.rectangle((x, y, x + 50, y + 50), fill=rnd.choice([NAVY, TEAL, ORANGE]))
    im.save(path, optimize=True)


def pattern(path):
    """Dotted pattern laid over the header band (semi-transparent)."""
    w, h = 2400, 160
    im = Image.new("RGBA", (w, h), (0, 0, 0, 0))
    d = ImageDraw.Draw(im)
    for y in range(12, h, 24):
        for x in range(12 + (y // 24 % 2) * 12, w, 24):
            d.ellipse((x - 3, y - 3, x + 3, y + 3), fill=(255, 255, 255, 60))
    im.save(path, optimize=True)


def icon(path, size=200):
    """Small rounded-square mark for the footer."""
    im = Image.new("RGBA", (size, size), (0, 0, 0, 0))
    d = ImageDraw.Draw(im)
    d.rounded_rectangle((10, 10, size - 10, size - 10), radius=40, fill=TEAL + (255,))
    for i, y in enumerate((60, 95, 130)):
        d.rectangle((50, y, size - 50 - i * 25, y + 16), fill=(255, 255, 255, 255))
    im.save(path, optimize=True)


# ---- pptx ----------------------------------------------------------------------------------------------------------

_next_id = [900]
def next_id():
    _next_id[0] += 1
    return _next_id[0]


def add_picture(part, sp_tree, image, left, top, width, height, name="Picture", index=None):
    """A <p:pic> on a master / layout (python-pptx offers add_picture only for slides)."""
    _, rId = part.get_or_add_image_part(str(image))
    pic = CT_Picture.new_pic(next_id(), name, name, rId, left, top, width, height)
    if index is None:
        sp_tree.append(pic)
    else:
        sp_tree.insert(index, pic)
    return pic


def add_rect(shapes, sp_tree, left, top, width, height, theme_color=None, rgb=None, name="Rect", index=None):
    sp = CT_Shape.new_autoshape_sp(next_id(), name, "rect", left, top, width, height)
    if index is None:
        sp_tree.append(sp)
    else:
        sp_tree.insert(index, sp)
    shape = Shape(sp, shapes)
    shape.fill.solid()
    if theme_color is not None:
        shape.fill.fore_color.theme_color = theme_color
    else:
        shape.fill.fore_color.rgb = RGBColor(*rgb)
    shape.line.fill.background()
    return shape


def set_lvl1(placeholder, rgb=None, size_pt=None, align=None, bold=None):
    """First-level default run properties on a placeholder (what the preview reads as its style)."""
    lst = placeholder._element.txBody.find(qn("a:lstStyle"))
    if lst is None:
        lst = etree.SubElement(placeholder._element.txBody, qn("a:lstStyle"))
    lvl = parse_xml(f'<a:lvl1pPr {nsdecls("a")}><a:defRPr/></a:lvl1pPr>')
    if align:
        lvl.set("algn", align)
    rpr = lvl.find(qn("a:defRPr"))
    if size_pt:
        rpr.set("sz", str(int(size_pt * 100)))
    if bold is not None:
        rpr.set("b", "1" if bold else "0")
    if rgb:
        rpr.append(parse_xml(f'<a:solidFill {nsdecls("a")}><a:srgbClr val="{"%02X%02X%02X" % rgb}"/></a:solidFill>'))
    lst.insert(0, lvl)


def set_theme(master):
    theme_part = master.part.part_related_by(RT.THEME)
    root = etree.fromstring(theme_part.blob)
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    for scope in ("majorFont", "minorFont"):
        el = root.find(f".//a:{scope}", ns)
        el.find("a:latin", ns).set("typeface", FONT_LATIN)
        for f in el.findall("a:font", ns):
            if f.get("script") == "Jpan":
                f.set("typeface", FONT_JA)
    scheme = root.find(".//a:clrScheme", ns)
    for key, rgb in {"dk2": INK, "lt2": PAPER, "accent1": NAVY, "accent2": TEAL, "accent3": ORANGE, "accent4": (122, 90, 248), "accent5": (46, 204, 113), "accent6": (230, 57, 70)}.items():
        node = scheme.find(f"a:{key}", ns)
        for child in list(node):
            node.remove(child)
        etree.SubElement(node, qn("a:srgbClr")).set("val", "%02X%02X%02X" % rgb)
    theme_part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def placeholders_of(shapes, ph_type):
    return [s for s in shapes if s.is_placeholder and s.placeholder_format.type == ph_type]


def main():
    tmp = Path(tempfile.mkdtemp(prefix="mdslide-master-"))
    imgs = {k: tmp / f"{k}.png" for k in ("hero", "logo", "mountains", "waves", "grid", "pattern", "icon")}
    hero(imgs["hero"]); logo(imgs["logo"]); pattern(imgs["pattern"]); icon(imgs["icon"])
    for k in ("mountains", "waves", "grid"):
        tile(imgs[k], k)

    prs = Presentation(str(SRC))
    master = prs.slide_master
    set_theme(master)
    W, H = prs.slide_width, prs.slide_height

    # Master: header band (with pattern and logo) and footer band, behind the placeholders (inserted at the front).
    tree = master.shapes._spTree
    band_h, foot_h = Inches(0.55), Inches(0.45)
    add_rect(master.shapes, tree, 0, 0, W, band_h, theme_color=MSO_THEME_COLOR.ACCENT_1, name="Header band", index=2)          # 5 = accent1
    add_picture(master.part, tree, imgs["pattern"], 0, 0, W, band_h, name="Header pattern", index=3)
    add_picture(master.part, tree, imgs["logo"], W - Inches(0.9), Inches(0.1), Inches(0.35), Inches(0.35), name="Logo small", index=4)
    add_rect(master.shapes, tree, 0, H - foot_h, W, foot_h, rgb=PAPER, name="Footer band", index=5)
    add_picture(master.part, tree, imgs["icon"], Inches(0.45), H - foot_h + Inches(0.08), Inches(0.29), Inches(0.29), name="Footer icon", index=6)
    for ph in placeholders_of(master.shapes, PP_PLACEHOLDER.TITLE):
        ph.top, ph.height = Inches(0.75), Inches(1.1)
    for ph in placeholders_of(master.shapes, PP_PLACEHOLDER.BODY):
        ph.top, ph.height = Inches(1.95), H - Inches(1.95) - foot_h - Inches(0.15)
    for ph in placeholders_of(master.shapes, PP_PLACEHOLDER.FOOTER):
        ph.left, ph.top, ph.width, ph.height = Inches(0.85), H - foot_h, Inches(7), foot_h
        ph.text_frame.text = "ACME Platform Engineering · Confidential"
        ph.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = ph.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.LEFT
        p.runs[0].font.size = Pt(9); p.runs[0].font.color.rgb = RGBColor(*INK)
    for ph in placeholders_of(master.shapes, PP_PLACEHOLDER.DATE):
        ph.left, ph.top, ph.width, ph.height = Inches(8.2), H - foot_h, Inches(3.2), foot_h
        set_lvl1(ph, rgb=(110, 120, 140), size_pt=9, align="r")
    for ph in placeholders_of(master.shapes, PP_PLACEHOLDER.SLIDE_NUMBER):
        ph.left, ph.top, ph.width, ph.height = Inches(11.6), H - foot_h, Inches(1.3), foot_h
        set_lvl1(ph, rgb=(110, 120, 140), size_pt=9, align="r")

    for layout in prs.slide_layouts:
        ltree = layout.shapes._spTree
        if layout.name == "Cover":
            layout._element.set("showMasterSp", "0")
            _, rId = layout.part.get_or_add_image_part(str(imgs["hero"]))
            bg = parse_xml(f'<p:bg {nsdecls("p", "a", "r")}><p:bgPr><a:blipFill dpi="0" rotWithShape="1"><a:blip r:embed="{rId}"/><a:srcRect/><a:stretch><a:fillRect/></a:stretch></a:blipFill><a:effectLst/></p:bgPr></p:bg>')
            layout._element.cSld.insert(0, bg)
            add_picture(layout.part, ltree, imgs["logo"], Inches(0.6), Inches(0.5), Inches(1.3), Inches(1.3), name="Logo", index=2)
            x = W - Inches(4.4)
            for i, k in enumerate(("mountains", "waves", "grid")):
                add_picture(layout.part, ltree, imgs[k], x, Inches(0.5) + i * Inches(2.25), Inches(3.9), Inches(2.05), name=f"Tile {k}", index=3 + i)
            for ph in placeholders_of(layout.shapes, PP_PLACEHOLDER.CENTER_TITLE):
                ph.left, ph.top, ph.width, ph.height = Inches(0.6), Inches(4.3), Inches(7.8), Inches(1.5)
                set_lvl1(ph, rgb=(255, 255, 255), size_pt=44, align="l", bold=True)
                ph.text_frame.vertical_anchor = MSO_ANCHOR.BOTTOM
            for ph in placeholders_of(layout.shapes, PP_PLACEHOLDER.SUBTITLE):
                ph.left, ph.top, ph.width, ph.height = Inches(0.6), Inches(5.85), Inches(7.8), Inches(1.1)
                set_lvl1(ph, rgb=(220, 230, 240), size_pt=18, align="l")
                ph.text_frame.vertical_anchor = MSO_ANCHOR.TOP
        elif layout.name == "Section":
            add_picture(layout.part, ltree, imgs["waves"], 0, band_h, Inches(4.2), H - band_h - foot_h, name="Side band", index=2)
            for ph in placeholders_of(layout.shapes, PP_PLACEHOLDER.TITLE):
                ph.left, ph.top, ph.width, ph.height = Inches(4.8), Inches(2.6), Inches(8), Inches(1.6)
                set_lvl1(ph, rgb=NAVY, size_pt=36, align="l", bold=True)
            for ph in placeholders_of(layout.shapes, PP_PLACEHOLDER.BODY):
                ph.left, ph.top, ph.width, ph.height = Inches(4.8), Inches(4.3), Inches(8), Inches(1.2)
                set_lvl1(ph, rgb=(110, 120, 140), size_pt=16, align="l")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT))
    print(f"wrote {OUT} ({OUT.stat().st_size // 1024} KB)")


if __name__ == "__main__":
    main()
