"""Real Electron app: folder from argv, master auto-import, paste to disk, chokidar reload, Python export."""
import pytest
from pptx import Presentation

from conftest import PASTE_JS, png_b64

pytestmark = pytest.mark.electron


def dismiss_help(pg):
    pg.get_by_role("banner").wait_for(timeout=10000)  # the app renders after settings.json is read
    if pg.get_by_text("mdslide の使い方").is_visible():
        pg.get_by_role("button", name="閉じる").click()


def wait_js(pg, expression, timeout_ms=10000):
    """Poll a page condition from Python: the app's CSP forbids eval, which Playwright's in-page waiting uses (ADR-0026)."""
    for _ in range(max(1, timeout_ms // 100)):
        if pg.evaluate(expression):
            return
        pg.wait_for_timeout(100)
    raise AssertionError(f"timed out after {timeout_ms} ms: {expression[:100]}")


def wait_until(pg, pred, timeout_ms=15000):
    """Poll a Python-side condition (disk state) instead of sleeping a fixed time: CI runners are slower than a laptop."""
    for _ in range(max(1, timeout_ms // 250)):
        if pred():
            return True
        pg.wait_for_timeout(250)
    return pred()


def deck_text(ws):
    return (ws / "deck.md").read_text(encoding="utf8") if (ws / "deck.md").exists() else ""


def test_workspace_from_argv_paste_watch_and_pptx_export(electron_app):
    pg, ws = electron_app
    dismiss_help(pg)
    assert pg.get_by_role("button", name="ws/deck.md").is_visible()  # toolbar shows folder/file
    # the launch check finds a Python that can import python-pptx (the export below uses the same one)
    check = pg.evaluate("() => window.mdslide.checkPython()")
    assert check["ok"] and check["pptx"], check
    assert pg.get_by_text("python-pptx が入っていません").count() == 0
    assert (ws / "deck.md").exists()
    assert pg.get_by_role("banner").get_by_role("combobox").input_value() == "ws:ws"

    pg.locator(".cm-content").click(); pg.keyboard.press("Escape"); pg.keyboard.type("29G")
    pg.evaluate(PASTE_JS, png_b64(600, 900))
    # the image is written first, then the reference is autosaved into deck.md (1.5 s debounce)
    assert wait_until(pg, lambda: (ws / "images").exists() and [p.name for p in (ws / "images").iterdir()] == ["2-1-パイプラインの構成.png"])
    assert wait_until(pg, lambda: "images/2-1-パイプラインの構成.png" in deck_text(ws))
    pg.get_by_role("button", name="保存済み").wait_for(timeout=10000)

    # chokidar picks up an external rewrite
    t = deck_text(ws)
    (ws / "deck.md").write_text(t.replace("# 背景と目的", "# 背景と目的（外部更新）"), encoding="utf8")
    pg.locator(".nav-item .label", has_text="1. 背景と目的（外部更新）").wait_for(timeout=10000)
    assert pg.locator(".nav-item .label", has_text="1. 背景と目的（外部更新）").count() == 1

    # switch the pasted image to 3/4 left, then export a pptx through the bundled Python tool
    pg.locator(".nav-item", has_text="2.1.").first.click()
    pg.get_by_role("button", name="3/4").click(); pg.get_by_role("button", name="左").click()
    assert wait_until(pg, lambda: "{img=3/4 side=left}" in deck_text(ws))
    pg.get_by_role("button", name="保存済み").wait_for(timeout=10000)
    pg.get_by_role("button", name="書き出す").click()
    pg.get_by_text("out/deck.pptx を生成しました").wait_for(timeout=30000)
    out = ws / "out" / "deck.pptx"
    assert out.exists()
    prs = Presentation(str(out))
    titles = [s.shapes.title.text for s in prs.slides if s.shapes.title is not None]
    assert "2.1. パイプラインの構成" in titles
    img_slide = [s for s in prs.slides if s.shapes.title is not None and s.shapes.title.text == "2.1. パイプラインの構成"][0]
    pics = [sh for sh in img_slide.shapes if sh.shape_type == 13]
    assert len(pics) == 1 and abs(pics[0].width / pics[0].height - 600 / 900) < 0.01
    assert pics[0].left < prs.slide_width / 2  # side=left

    # ADR-0026: the file IPC reaches only the opened folders, the page has a CSP, and it cannot open windows or navigate away
    outside = ws.parent / "outside.txt"
    outside.write_text("secret", encoding="utf8")
    assert pg.evaluate("(p) => window.mdslide.readText(p)", str(outside)) is None
    refused = pg.evaluate("(p) => window.mdslide.writeText(p, 'x').then(() => 'written', (e) => String(e))", str(ws.parent / "evil.txt"))
    assert "開いたフォルダの外" in refused and not (ws.parent / "evil.txt").exists()
    assert pg.evaluate("(p) => window.mdslide.readText(p)", str(ws / "images" / ".." / ".." / "outside.txt")) is None
    assert "script-src 'self'" in pg.evaluate("() => document.querySelector('meta[http-equiv=\"Content-Security-Policy\"]').content")
    assert pg.evaluate("() => window.open('https://example.com/') === null")
    url = pg.url
    pg.evaluate("() => { location.href = 'https://example.com/'; }")
    pg.wait_for_timeout(800)
    assert pg.url == url
    assert pg.errors == []


TERM_TEXT = "()=>{const t=window.__mdslideTerminal;const b=t.buffer.active;const o=[];for(let i=0;i<b.length;i++)o.push(b.getLine(i)?.translateToString(true)??'');return o.join(' ');}"


def test_console_is_a_real_terminal_that_starts_claude(electron_app):
    pg, ws = electron_app
    pg.get_by_test_id("console").wait_for(timeout=10000)
    dismiss_help(pg)
    assert "deck.md" in (ws / "AGENTS.md").read_text(encoding="utf8")           # conventions for any agent
    assert (ws / "CLAUDE.md").read_text(encoding="utf8") == "@AGENTS.md\n"       # Claude Code imports the same file
    # the shell starts in the deck folder and `claude` is launched automatically
    wait_js(pg, f"({TERM_TEXT})().includes('FAKE CLAUDE READY')", 15000)
    pg.locator(".terminal-host").click()
    pg.keyboard.type("Risk list"); pg.keyboard.press("Enter")
    wait_js(pg, f"({TERM_TEXT})().includes('claude got: Risk list')", 10000)
    assert "## Risk list" in (ws / "deck.md").read_text(encoding="utf8")
    pg.wait_for_timeout(1500)  # chokidar -> reload
    assert pg.locator(".nav-item .label", has_text="Risk list").count() == 1
    # leave the fake claude (EOF ends its read loop) and check from the shell that cwd is the workspace.
    # The marker is computed by the shell so a long, wrapped path in the terminal buffer cannot break the check.
    pg.keyboard.press("Control+d")
    pg.wait_for_timeout(500)
    pg.keyboard.type(f'echo CWD-$([ "$PWD" = "{ws}" ] && echo MATCH || echo MISMATCH)'); pg.keyboard.press("Enter")
    wait_js(pg, f"({TERM_TEXT})().includes('CWD-MATCH')", 10000)
    pg.screenshot(path=str(ws.parent / "console.png"))  # kept as a test artifact
    # tool settings: add a custom tool, select it, launch it from the bar
    pg.get_by_role("button", name="ツール設定").click()
    pg.get_by_role("textbox", name="新しいツールの名前").fill("Echo")
    pg.get_by_role("textbox", name="新しいツールのコマンド").fill("echo")
    pg.get_by_role("textbox", name="新しいツールの引数").fill("CUSTOM-TOOL-OK")
    pg.get_by_role("button", name="追加").click()
    pg.get_by_role("radio", name="Echo を選択").check()
    pg.get_by_role("button", name="閉じる").click()
    pg.get_by_role("button", name="起動").click()
    wait_js(pg, f"({TERM_TEXT})().includes('CUSTOM-TOOL-OK')", 10000)
    pg.get_by_role("combobox", name="起動するツール").select_option("gemini")
    # settings live in a JSON file (MDSLIDE_CONFIG), written atomically after a short debounce
    import json, time
    cfg = ws.parent / "config" / "settings.json"
    for _ in range(20):
        if cfg.exists() and json.loads(cfg.read_text(encoding="utf8")).get("tools", {}).get("selectedId") == "gemini":
            break
        time.sleep(0.25)
    data = json.loads(cfg.read_text(encoding="utf8"))
    assert data["tools"]["selectedId"] == "gemini"
    assert data["workspace"]["lastPath"] == str(ws)
    assert [t["id"] for t in data["tools"]["items"] if t["command"] == "echo"]  # the custom tool
    assert [t["id"] for t in data["tools"]["items"] if t["id"] in {"claude", "codex", "gemini"}] == []  # untouched presets are not written
    # a hand edit is picked up when the window regains focus
    data["tools"]["selectedId"] = "aider"
    cfg.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf8")
    pg.evaluate("()=>window.dispatchEvent(new Event('focus'))")
    wait_js(pg, "()=>document.querySelector('select[aria-label=\"起動するツール\"]').value==='aider'", 5000)
    # ⌘J hides and shows the console
    pg.keyboard.press("Meta+J"); pg.wait_for_timeout(200)
    assert pg.get_by_test_id("console").count() == 0
    pg.keyboard.press("Meta+J"); pg.wait_for_timeout(200)
    assert pg.get_by_test_id("console").count() == 1
    # dragging the console taller shrinks the slide; it must not spill out of the preview pane
    sep = pg.get_by_role("separator", name="コンソールの高さ").bounding_box()
    pg.mouse.move(sep["x"] + sep["width"] / 2, sep["y"] + 2); pg.mouse.down()
    pg.mouse.move(sep["x"] + sep["width"] / 2, sep["y"] - 260, steps=5); pg.mouse.up()
    pg.wait_for_timeout(300)
    fit = pg.evaluate("()=>{const f=document.querySelector('.slide-frame').getBoundingClientRect(),p=document.querySelector('.slide-fit').getBoundingClientRect();return {ok:f.top>=p.top-1&&f.bottom<=p.bottom+1,w:f.width}}")
    assert fit["ok"] and fit["w"] > 100, fit
    # ⌘, opens the settings sheet; the appearance choice reaches nativeTheme, so prefers-color-scheme follows it (theme:set IPC).
    # Playwright emulates a light scheme on every page it attaches to; drop that so the query reports what Electron decides.
    pg.emulate_media(color_scheme="no-override")
    pg.keyboard.press("Meta+,")
    pg.get_by_role("dialog", name="設定").wait_for(timeout=5000)
    pg.get_by_role("button", name="ダーク").click()
    wait_js(pg, "()=>matchMedia('(prefers-color-scheme: dark)').matches && document.documentElement.dataset.theme==='dark'", 5000)
    pg.get_by_role("button", name="ライト").click()
    wait_js(pg, "()=>!matchMedia('(prefers-color-scheme: dark)').matches && document.documentElement.dataset.theme==='light'", 5000)
    pg.get_by_role("button", name="自動").click()
    wait_js(pg, "()=>document.documentElement.dataset.theme===undefined", 5000)
    wait_until(pg, lambda: json.loads(cfg.read_text(encoding="utf8")).get("appearance", {}).get("theme") == "auto")
    pg.get_by_role("button", name="閉じる").click()
    assert pg.get_by_role("dialog", name="設定").count() == 0
    assert pg.errors == []


def test_inbox_to_deck_with_figures_and_undo(electron_app):
    pg, ws = electron_app
    dismiss_help(pg)
    assert (ws / "theme.json").exists() and (ws / "tools" / "mdslide_draw.py").exists()
    wait_js(pg, f"({TERM_TEXT})().includes('FAKE CLAUDE READY')", 15000)
    before = (ws / "deck.md").read_text(encoding="utf8")
    pg.get_by_role("button", name="下書き").click()
    box = pg.get_by_role("textbox", name="メモ")
    box.click(); box.fill("今期はデプロイ頻度を上げたい。理由は障害対応が属人化しているから。")
    box.blur()
    wait_js(pg, "()=>document.querySelectorAll('.inbox-notes .name').length>=1", 5000)
    assert [p.suffix for p in (ws / "notes").iterdir() if p.suffix] == [".md"]
    pg.get_by_role("button", name="整形して deck.md に").click()
    wait_js(pg, f"({TERM_TEXT})().includes('claude got:')", 10000)
    wait_js(pg, "()=>[...document.querySelectorAll('.nav-item .label')].some(e=>e.textContent.includes('メモから'))", 15000)
    import time as _t
    for _ in range(60):  # the fake tool draws the figure after rewriting deck.md
        text = (ws / "deck.md").read_text(encoding="utf8")
        if "![流れ](images/flow.png)" in text:
            break
        _t.sleep(0.5)
    assert "- 今期はデプロイ頻度を上げたい" in text and "![流れ](images/flow.png)" in text
    assert (ws / "images" / "flow.png").exists()
    # the preview shows the generated figure
    pg.locator(".nav-item", has_text="メモから").first.click()
    wait_js(pg, "()=>{const i=document.querySelector('.slide-frame img');return i&&i.naturalWidth>0}", 10000)
    pg.screenshot(path=str(ws.parent / "inbox.png"))
    # undo restores the previous deck
    pg.get_by_role("button", name="前の版に戻す").click()
    wait_js(pg, "()=>![...document.querySelectorAll('.nav-item .label')].some(e=>e.textContent.includes('メモから'))", 10000)
    assert (ws / "deck.md").read_text(encoding="utf8") == before
    assert pg.errors == []
