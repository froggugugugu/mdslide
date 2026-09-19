"""tools/convert_master.py: a template becomes an mdslide master without losing the design it came with."""
import subprocess
import sys

import pytest
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Emu

from conftest import ROOT, slide

SCRIPT = ROOT / "tools" / "convert_master.py"
W, H = 12192000, 6858000
PH_PATH = f'{qn("p:nvSpPr")}/{qn("p:nvPr")}/{qn("p:ph")}'
ROLES = ("Cover", "Agenda", "Section", "Body-Text", "Body-2col")


def kind(shape):
    el = shape._element.find(PH_PATH)
    return None if el is None else (el.get("type") or "obj")


def bodies(layout):
    return [s for s in layout.placeholders if kind(s) in ("body", "obj")]


def convert(src, out):
    """Run the converter in-process (so coverage sees it), mimicking the CLI contract."""
    import contextlib
    import io

    import convert_master as m

    argv = sys.argv
    sys.argv = ["convert_master.py", str(src), "-o", str(out)]
    try:
        with contextlib.redirect_stdout(io.StringIO()) as so:
            m.main()
    finally:
        sys.argv = argv
    return Presentation(str(out)), so.getvalue()


@pytest.fixture
def office(tmp_path):
    """A plain Office template: none of the mdslide names, a cramped content box, and leftover slides."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    body = next(p for p in prs.slide_layouts[1].placeholders if p.placeholder_format.idx == 1)
    body.left, body.top, body.width, body.height = Emu(700000), Emu(1800000), Emu(1800000), Emu(500000)
    prs.slides.add_slide(prs.slide_layouts[0])
    prs.slides.add_slide(prs.slide_layouts[1])
    p = tmp_path / "office.pptx"
    prs.save(str(p))
    return p


@pytest.fixture
def corporate(tmp_path):
    """A decorated template under corporate names, with no layout that could serve as the agenda."""
    prs = Presentation(str(ROOT / "examples" / "decorated-master.pptx"))
    prs.slide_layouts.remove(prs.slide_layouts.get_by_name("Agenda"))
    for old, new in (("Cover", "表紙 2026"), ("Body-Text", "本文"), ("Body-2col", "比較"), ("Section", "扉")):
        prs.slide_layouts.get_by_name(old).name = new
    p = tmp_path / "corporate.pptx"
    prs.save(str(p))
    return p


def test_cli_entrypoint(tmp_path, office):
    out = tmp_path / "master.pptx"
    r = subprocess.run([sys.executable, str(SCRIPT), str(office), "-o", str(out)], capture_output=True, text=True)
    assert r.returncode == 0, r.stderr
    assert out.exists() and "Body-Text" in r.stdout


def test_missing_source_is_reported(tmp_path):
    r = subprocess.run([sys.executable, str(SCRIPT), str(tmp_path / "nope.pptx"), "-o", str(tmp_path / "m.pptx")], capture_output=True, text=True)
    assert r.returncode != 0 and "ありません" in r.stderr


def test_every_role_is_named_and_the_deck_is_emptied(tmp_path, office):
    prs, report = convert(office, tmp_path / "master.pptx")
    assert set(ROLES) <= {l.name for l in prs.slide_layouts}
    assert len(prs.slides._sldIdLst) == 0  # a master carries no pages
    assert all(role in report for role in ROLES)


def test_body_box_is_resized_so_a_slide_holds_a_paragraph(tmp_path, office):
    """The reported failure: a cramped source box made mdslide split every line onto its own slide."""
    before = next(p for p in Presentation(str(office)).slide_layouts[1].placeholders if p.placeholder_format.idx == 1)
    assert before.width < 0.2 * W  # the source box really is narrow
    prs, _ = convert(office, tmp_path / "master.pptx")
    body = bodies(prs.slide_layouts.get_by_name("Body-Text"))[0]
    assert body.width >= 0.85 * W
    assert body.height / 12700 / (18 * 1.3) >= 8  # at least eight 18pt lines fit


def test_body_box_keeps_clear_of_the_header_and_footer(tmp_path, corporate):
    prs, _ = convert(corporate, tmp_path / "master.pptx")
    master = prs.slide_master
    header = max(s.top + s.height for s in master.shapes if not s.is_placeholder and s.top + s.height <= H / 3)
    footer = min(s.top for s in master.shapes if not s.is_placeholder and s.top >= H * 2 / 3)
    for role in ("Body-Text", "Body-2col", "Agenda"):
        for body in bodies(prs.slide_layouts.get_by_name(role)):
            assert body.top >= header, role
            assert body.top + body.height <= footer, role


def test_two_columns_are_side_by_side(tmp_path, corporate):
    prs, _ = convert(corporate, tmp_path / "master.pptx")
    left, right = sorted(bodies(prs.slide_layouts.get_by_name("Body-2col")), key=lambda s: s.left)
    assert left.left + left.width <= right.left
    assert abs(left.width - right.width) <= 1


def test_a_missing_role_is_cloned_with_its_decoration(tmp_path, corporate):
    prs, report = convert(corporate, tmp_path / "master.pptx")
    agenda = prs.slide_layouts.get_by_name("Agenda")
    assert len(bodies(agenda)) == 1
    assert [s for s in agenda.placeholders if kind(s) == "title"]
    assert "複製" in report


def test_theme_decoration_and_furniture_survive(tmp_path, corporate):
    src = Presentation(str(corporate))
    prs, _ = convert(corporate, tmp_path / "master.pptx")
    assert prs.slide_master.part.part_related_by(
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/theme").blob == src.slide_master.part.part_related_by(
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/theme").blob
    pictures = [s.name for s in prs.slide_master.shapes if not s.is_placeholder]
    assert len(pictures) == len([s.name for s in src.slide_master.shapes if not s.is_placeholder])
    furniture = {kind(s) for s in prs.slide_layouts.get_by_name("Body-Text").placeholders}
    assert {"dt", "ftr", "sldNum"} <= furniture


def test_extra_content_boxes_are_removed(tmp_path, office):
    prs, _ = convert(office, tmp_path / "master.pptx")
    for role in ("Body-Text", "Agenda"):
        assert len(bodies(prs.slide_layouts.get_by_name(role))) == 1, role
    assert len(bodies(prs.slide_layouts.get_by_name("Body-2col"))) == 2
    assert not bodies(prs.slide_layouts.get_by_name("Section"))


def test_names_that_would_shadow_a_role_are_moved_aside(tmp_path, office):
    prs, _ = convert(office, tmp_path / "master.pptx")
    import convert_master as m

    for layout in prs.slide_layouts:
        if layout.name not in ROLES:
            assert m.reads_as_role(layout.name) is None, layout.name


@pytest.mark.parametrize("name, role", [
    ("Cover", "Cover"), ("表紙", "Cover"), ("Title Slide", "Cover"),
    ("AGENDA", "Agenda"), ("目次", "Agenda"),
    ("Section Header", "Section"), ("中表紙", "Section"),
    ("Body-Text", "Body-Text"), ("body text", "Body-Text"),
    ("Body-2col", "Body-2col"), ("Body-TwoCol", "Body-2col"),
    ("Body-Weird", None), ("Title Only", None), ("Title and Vertical Text", None), ("Comparison", None),
])
def test_role_names_are_read_like_the_app_reads_them(name, role):
    """The table of roleFromLayoutName in tests/unit/importMaster.test.ts: the mirror must not drift from it."""
    import convert_master as m

    assert m.reads_as_role(name) == role


def test_a_converted_master_exports(tmp_path, office, deck_json):
    """The point of converting: the result carries a deck through tools/export_pptx.py."""
    master = tmp_path / "master.pptx"
    convert(office, master)
    deck = deck_json([slide("cover", "T"), slide("section", "S"), slide("body", "B", body=["- a", "- b"])])
    out = tmp_path / "out.pptx"
    r = subprocess.run([sys.executable, str(ROOT / "tools" / "export_pptx.py"), str(deck), "--master", str(master), "-o", str(out)],
                       capture_output=True, text=True)
    assert r.returncode == 0, r.stderr
    assert "no body placeholder" not in r.stderr
    exported = Presentation(str(out))
    assert len(exported.slides) == 3
    assert "B" in [s.text for s in exported.slides[2].shapes if s.has_text_frame]


def test_an_mdslide_master_converts_to_itself(tmp_path):
    """Running it on a master that already follows the convention keeps the roles and stays importable."""
    prs, _ = convert(ROOT / "examples" / "sample-master.pptx", tmp_path / "master.pptx")
    assert {r for r in ROLES} <= {l.name for l in prs.slide_layouts}
    assert len(bodies(prs.slide_layouts.get_by_name("Body-Text"))) == 1
    assert len(prs.slides._sldIdLst) == 0
