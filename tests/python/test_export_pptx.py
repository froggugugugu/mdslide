import subprocess
import sys
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Emu, Pt

from conftest import ROOT, geometry, slide

SCRIPT = ROOT / "tools" / "export_pptx.py"


class Result:
    def __init__(self, code, out, err):
        self.returncode, self.stdout, self.stderr = code, out, err


def run(deck, master, out, assets):
    """Run the exporter in-process (so coverage sees it), mimicking the CLI contract."""
    import io, contextlib
    import export_pptx as m
    argv = sys.argv
    sys.argv = ["export_pptx.py", str(deck), "--master", str(master), "-o", str(out), "--assets", str(assets)]
    so, se = io.StringIO(), io.StringIO()
    code = 0
    try:
        with contextlib.redirect_stdout(so), contextlib.redirect_stderr(se):
            m.main()
    except SystemExit as e:
        code = 1 if isinstance(e.code, str) else (e.code or 0)
        if isinstance(e.code, str):
            se.write(e.code)
    finally:
        sys.argv = argv
    return Result(code, so.getvalue(), se.getvalue())


def test_cli_entrypoint(tmp_path, master, deck_json):
    r = subprocess.run([sys.executable, str(SCRIPT), str(deck_json([slide("body", "B", body=["x"])])), "--master", str(master), "-o", str(tmp_path / "cli.pptx")],
                       capture_output=True, text=True)
    assert r.returncode == 0 and (tmp_path / "cli.pptx").exists()


def test_full_deck_roundtrip(tmp_path, master, deck_json, png):
    png("fig.png", 400, 800)
    slides = [
        slide("cover", "Deck", master_layout="Cover", body=["Sub"]),
        slide("agenda", "Agenda", master_layout="Agenda", agenda={"items": [{"number": "1", "title": "A"}, {"number": "2", "title": "B"}]}),
        slide("section", "1. A", master_layout="Section"),
        slide("body", "1.1. Text", master_layout="Body-Text", body=["- **bold** x", "  - nested", "para `code`"], notes=["speaker note"]),
        slide("body", "1.2. Table", master_layout="Body-Text", body=["intro", "", "| k | v |", "|---|---|", "| a | b |"]),
        slide("body", "1.3. Two", layout="2col", master_layout="Body-2col", body=["- l", "", "- r"]),
        slide("body", "2.1. Img", layout="image", master_layout="Body-Text", body=["- t"], images=[{"alt": "f", "src": "images/fig.png"}], geometry=geometry("right", 0.5)),
    ]
    out = tmp_path / "out" / "deck.pptx"   # out/ does not exist yet
    r = run(deck_json(slides), master, out, tmp_path)
    assert r.returncode == 0, r.stderr
    assert "7 slides" in r.stdout
    prs = Presentation(str(out))
    assert len(prs.slides) == 7
    s = list(prs.slides)
    assert s[0].slide_layout.name == "Cover" and s[0].shapes.title.text == "Deck"
    agenda_text = "\n".join(sh.text_frame.text for sh in s[1].shapes if sh.has_text_frame)
    assert "1. A" in agenda_text and "2. B" in agenda_text
    # rich text
    body = [sh for sh in s[3].placeholders if sh.placeholder_format.idx != 0][0]
    runs = [r for p in body.text_frame.paragraphs for r in p.runs]
    assert any(r.font.bold for r in runs) and any(r.font.name == "Consolas" for r in runs)
    assert body.text_frame.paragraphs[1].level == 1
    assert s[3].notes_slide.notes_text_frame.text == "speaker note"
    # native table
    tables = [sh for sh in s[4].shapes if sh.has_table]
    assert len(tables) == 1 and tables[0].table.cell(0, 0).text == "k" and tables[0].table.cell(1, 1).text == "b"
    # two columns
    cols = [sh.text_frame.text for sh in s[5].placeholders if sh.placeholder_format.idx != 0 and sh.placeholder_format.type not in (13, 15, 16)]  # not the slide number / footer / date
    assert cols == ["l", "r"]
    # image slide: picture fitted (aspect 1:2) inside the box, anchored right; body resized to the left area
    g = geometry("right", 0.5)
    pics = [sh for sh in s[6].shapes if sh.shape_type == 13]
    assert len(pics) == 1
    p = pics[0]
    assert abs(p.width / p.height - 0.5) < 0.01
    assert p.top == g["imageBox"]["y"]
    assert p.left + p.width == g["imageBox"]["x"] + g["imageBox"]["w"]
    assert p.height <= g["imageBox"]["h"]
    body6 = [sh for sh in s[6].placeholders if sh.placeholder_format.idx != 0][0]
    assert body6.left == g["body"]["x"] and body6.width == g["body"]["w"]
    # empty placeholders were removed
    for sl in s:
        for sh in sl.placeholders:
            if sh.has_text_frame:
                assert sh.text_frame.text.strip() != ""


def test_full_width_image_puts_text_below_when_room(tmp_path, master, deck_json, png):
    png("wide.png", 1600, 300)
    png("tall.png", 300, 1600)
    slides = [
        slide("body", "W", layout="image", master_layout="Body-Text", body=["under"], images=[{"alt": "", "src": "images/wide.png"}], geometry=geometry("left", 1)),
        slide("body", "T", layout="image", master_layout="Body-Text", body=["no room"], images=[{"alt": "", "src": "images/tall.png"}], geometry=geometry("left", 1)),
    ]
    out = tmp_path / "o.pptx"
    r = run(deck_json(slides), master, out, tmp_path)
    assert r.returncode == 0, r.stderr
    s = list(Presentation(str(out)).slides)
    pic = [sh for sh in s[0].shapes if sh.shape_type == 13][0]
    body = [sh for sh in s[0].placeholders if sh.placeholder_format.idx != 0][0]
    assert body.top > pic.top + pic.height
    assert body.text_frame.text == "under"
    assert "no room for text" in r.stderr
    assert all(not (sh.has_text_frame and sh.text_frame.text == "no room") for sh in s[1].shapes)


def test_warnings_for_missing_and_placeholder_images(tmp_path, master, deck_json):
    slides = [
        slide("body", "M", layout="image", master_layout="Body-Text", images=[{"alt": "", "src": "images/none.png"}], geometry=geometry()),
        slide("body", "P", layout="image", master_layout="Body-Text", images=[{"alt": "TODO", "src": ""}], geometry=geometry()),
        slide("body", "X", master_layout="Body-Text", body=["t"], images=[{"alt": "", "src": "images/none.png"}]),
    ]
    r = run(deck_json(slides), master, tmp_path / "o.pptx", tmp_path)
    assert r.returncode == 0, r.stderr
    assert "image not found" in r.stderr
    assert "placeholder still empty" in r.stderr
    assert "non-image layout ignored" in r.stderr
    assert len(Presentation(str(tmp_path / "o.pptx")).slides) == 3


def test_layout_fallback_by_kind(tmp_path, master, deck_json):
    slides = [slide("cover", "C"), slide("section", "S"), slide("agenda", "A", agenda={"items": []}), slide("body", "B", body=["x"])]
    r = run(deck_json(slides), master, tmp_path / "o.pptx", tmp_path)
    assert r.returncode == 0, r.stderr
    names = [s.slide_layout.name for s in Presentation(str(tmp_path / "o.pptx")).slides]
    assert names == ["Cover", "Section", "Agenda", "Body-Text"]


def test_unknown_layout_and_bad_version_fail_clearly(tmp_path, master, deck_json):
    r = run(deck_json([slide("body", "B", master_layout="Body-Nope")]), master, tmp_path / "o.pptx", tmp_path)
    assert r.returncode == 0  # unknown name falls back to Body-Text by kind
    # a master without any body layout cannot be used
    prs = Presentation(str(master))
    for l in prs.slide_layouts:
        if l.name.lower().startswith("body"):
            l.name = "Other"
    nobody = tmp_path / "nobody.pptx"; prs.save(str(nobody))
    r = run(deck_json([slide("body", "B")]), nobody, tmp_path / "o2.pptx", tmp_path)
    assert r.returncode != 0 and "レイアウトが見つかりません" in r.stderr
    r = run(deck_json([], version=1), master, tmp_path / "o3.pptx", tmp_path)
    assert r.returncode != 0 and "version" in r.stderr


def test_fit_rect_and_helpers():
    import export_pptx as m
    box = {"x": 100, "y": 200, "w": 400, "h": 200}
    assert m.fit_rect(box, 1.0, "right") == (300, 200, 200, 200)
    assert m.fit_rect(box, 4.0, "left") == (100, 200, 400, 100)
    assert m.table_row("| a | b |") == ["a", "b"]
    assert m.table_row("plain") is None
    text, tables = m.split_tables(["t", "| a |", "|---|", "| b |", "u"])
    assert text == ["t", "u"] and tables == [[["a"], ["b"]]]
    assert m.split_columns(["a", "", "b"]) == (["a"], ["b"])
    assert m.split_columns(["a", "b", "c"]) == (["a", "b"], ["c"])
    assert m.image_size(Path("/nonexistent.png")) is None


def test_font_size_is_applied_and_overflow_is_reported(tmp_path, master, deck_json):
    slides = [
        slide("body", "Big", master_layout="Body-Text", body=["- a", "- b"], fontPt=28),
        slide("body", "Over", master_layout="Body-Text", body=[f"- 行{i}" for i in range(60)], fontPt=18),
    ]
    r = run(deck_json(slides), master, tmp_path / "o.pptx", tmp_path)
    assert r.returncode == 0, r.stderr
    s = list(Presentation(str(tmp_path / "o.pptx")).slides)
    body = [sh for sh in s[0].placeholders if sh.placeholder_format.idx != 0][0]
    assert all(run_.font.size == Pt(28) for p in body.text_frame.paragraphs for run_ in p.runs)
    assert "'Over' body is estimated at" in r.stderr and "18pt" in r.stderr


def test_fit_helpers_mirror_the_app():
    import export_pptx as m
    assert m.char_width_em("あ") == 1.0 and abs(m.char_width_em("a") - 0.55) < 1e-9
    assert m.display_lines("- " + "あ" * 30, 20) == 2
    assert m.display_lines("", 20) == 0.5
    assert m.display_lines("|---|", 20) == 0
    assert m.capacity_lines(288, 18) == 13


def test_footer_date_and_slide_number_placeholders_follow_the_layout(tmp_path, master, deck_json):
    """python-pptx clones no footer placeholders; the exporter copies them so slides carry the master's footer."""
    import json
    from pptx.enum.shapes import PP_PLACEHOLDER
    slides = [slide("cover", "Deck", master_layout="Cover"), slide("body", "1.1. A", master_layout="Body-Text", body=["x"]),
              slide("body", "1.2. B", master_layout="Body-Text", body=["y"])]
    p = deck_json(slides)
    d = json.loads(p.read_text(encoding="utf-8")); d["meta"]["date"] = "2026-09-08"; p.write_text(json.dumps(d), encoding="utf-8")
    out = tmp_path / "footer.pptx"
    r = run(p, master, out, tmp_path)
    assert r.returncode == 0, r.stderr
    prs = Presentation(str(out))
    for i, s in enumerate(prs.slides, start=1):
        by_type = {}
        for sh in s.placeholders:
            by_type.setdefault(sh.placeholder_format.type, []).append(sh)
        nums = by_type.get(PP_PLACEHOLDER.SLIDE_NUMBER, [])
        assert len(nums) == 1 and nums[0].text_frame.text == str(i)          # numbered, one per slide
        dates = by_type.get(PP_PLACEHOLDER.DATE, [])
        assert len(dates) == 1 and dates[0].text_frame.text == "2026-09-08"  # the deck's date as plain text
        assert not [r for r in dates[0]._element.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}fld")]  # no live field
        assert PP_PLACEHOLDER.FOOTER not in by_type                          # the layout's footer is empty: dropped
    # without a date in the deck, no date placeholder is added
    d["meta"].pop("date"); p.write_text(json.dumps(d), encoding="utf-8")
    out2 = tmp_path / "nodate.pptx"
    assert run(p, master, out2, tmp_path).returncode == 0
    for s in Presentation(str(out2)).slides:
        assert not [sh for sh in s.placeholders if sh.placeholder_format.type == PP_PLACEHOLDER.DATE]
        assert [sh for sh in s.placeholders if sh.placeholder_format.type == PP_PLACEHOLDER.SLIDE_NUMBER]


def test_decorated_master_exports_with_its_footer_and_pictures(tmp_path, deck_json):
    """examples/decorated-master.pptx (scripts/make_decorated_master.py): bands, pictures, a picture background and a fixed
    footer must come through the exporter untouched, and the copied footer carries the master's text."""
    from pptx.enum.shapes import PP_PLACEHOLDER
    master = ROOT / "examples" / "decorated-master.pptx"
    slides = [slide("cover", "Deck", master_layout="Cover", body=["Sub"]), slide("section", "1. A", master_layout="Section"),
              slide("body", "1.1. B", master_layout="Body-Text", body=["- x"])]
    out = tmp_path / "decorated.pptx"
    r = run(deck_json(slides), master, out, tmp_path)
    assert r.returncode == 0, r.stderr
    prs = Presentation(str(out))
    assert [s.slide_layout.name for s in prs.slides] == ["Cover", "Section", "Body-Text"]
    master_pics = [sh.name for sh in prs.slide_master.shapes if sh.shape_type == 13]
    assert set(master_pics) >= {"Header pattern", "Logo small", "Footer icon"}          # master pictures survive the round trip
    cover = prs.slides[0]
    assert cover.slide_layout._element.get("showMasterSp") == "0"
    assert [sh.name for sh in cover.slide_layout.shapes if sh.shape_type == 13] == ["Logo", "Tile mountains", "Tile waves", "Tile grid"]
    body = prs.slides[2]
    footers = [sh for sh in body.placeholders if sh.placeholder_format.type == PP_PLACEHOLDER.FOOTER]
    assert len(footers) == 1 and footers[0].text_frame.text == "ACME Platform Engineering · Confidential"
    nums = [sh for sh in body.placeholders if sh.placeholder_format.type == PP_PLACEHOLDER.SLIDE_NUMBER]
    assert len(nums) == 1 and nums[0].text_frame.text == "3"


def test_overflow_warning_uses_the_apps_estimate_when_deck_json_has_one(tmp_path, master, deck_json):
    """The app's estimate already counts the master's paragraph spacing, the insets and the footer; the exporter warns with it."""
    slides = [
        slide("body", "Tight", master_layout="Body-Text", body=["- a", "- b", "- c"], fontPt=18, fit={"used": 20.4, "capacity": 12}),
        slide("body", "Roomy", master_layout="Body-Text", body=[f"- 行{i}" for i in range(60)], fontPt=18, fit={"used": 10, "capacity": 12}),
    ]
    r = run(deck_json(slides), master, tmp_path / "warn.pptx", tmp_path)
    assert r.returncode == 0, r.stderr
    assert "'Tight' body is estimated at 20.4 lines but the box fits 12 at 18pt" in r.stderr
    assert "'Roomy'" not in r.stderr
